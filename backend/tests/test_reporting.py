import io
import re
import zipfile
from datetime import date
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.main import app
from app.services.downloads import report_content_disposition
from app.services.reporting import PDF_TYPE, SHEET_NAMES, XLSX_TYPE
from app.services.storage import get_forecast, save_forecast

client = TestClient(app)


@pytest.mark.parametrize("extension", ["pdf", "xlsx", "csv"])
def test_report_download_names_are_readable_and_browser_safe(extension):
    disposition = report_content_disposition(
        1, extension, start_date=date(2026, 2, 1), horizon_hours=48
    )
    assert (
        f'filename="WindFlow - Turbine 1 - Forecast from 2026-02-01 - 48h.{extension}"'
        in disposition
    )
    encoded_name = disposition.split("filename*=UTF-8''", 1)[1]
    assert (
        unquote(encoded_name) == f"WindFlow — Турбина 1 — Прогноз с 01.02.2026 — 48 ч.{extension}"
    )


def create_run(turbine=1, horizon=24, issue="2026-01-29"):
    response = client.post(
        "/api/forecasts",
        json={
            "turbine_id": turbine,
            "horizon_hours": horizon,
            "issue_date": issue,
            "data_mode": "offline",
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def pdf_text(response):
    assert response.status_code == 200, response.text[:200]
    assert response.headers["content-type"] == PDF_TYPE
    assert response.content.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(response.content))
    assert len(reader.pages) >= 2
    assert any(
        "/ToUnicode" in font.get_object()
        for page in reader.pages
        for font in page["/Resources"]["/Font"].values()
    )
    return "\n".join(page.extract_text() for page in reader.pages)


def xlsx_book(response):
    assert response.status_code == 200, response.text[:200]
    assert response.headers["content-type"] == XLSX_TYPE
    assert response.content.startswith(b"PK")
    assert "attachment" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == SHEET_NAMES
    return workbook


@pytest.mark.parametrize("turbine", [1, 2])
@pytest.mark.parametrize("horizon", [24, 48])
def test_real_run_exports_have_all_hours_and_russian_text(turbine, horizon):
    run = create_run(turbine, horizon)
    base = f"/api/forecasts/{run['run_id']}"
    text = pdf_text(client.get(base + "/report.pdf"))
    assert "Краткое резюме" in text
    assert "Почасовой прогноз" in text
    assert "windflow-v3" in text
    hourly_text = text.split("Почасовой прогноз", 1)[1]
    assert len(re.findall(r"\d{2}\.01\.2026 \d{2}:\d{2}", hourly_text)) == horizon
    workbook = xlsx_book(client.get(base + "/export.xlsx"))
    hourly = workbook["Hourly forecast"]
    assert hourly.max_row == horizon + 1
    assert hourly.cell(2, 1).value == run["forecast"][0]["timestamp"]
    assert hourly.cell(horizon + 1, 1).value == run["forecast"][-1]["timestamp"]
    assert hourly.cell(2, 4).value == pytest.approx(run["forecast"][0]["normalized_power"])
    assert workbook["Agent log"].max_row == len(run["steps"]) + 1
    assert workbook["Model metrics"].max_row == horizon // 24 + 1
    assert len(workbook["Summary"]._charts) == 1
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Нормированная выработка, турбино-часы"] == pytest.approx(
        run["expected_normalized_energy"]
    )
    assert sum(bool(row[8]) for row in hourly.iter_rows(min_row=2, values_only=True)) >= 1


@pytest.mark.parametrize("suffix", ["report.pdf", "export.xlsx"])
def test_unknown_or_incomplete_reports_are_not_downloaded(monkeypatch, february_weather, suffix):
    assert client.get(f"/api/forecasts/no-such-run/{suffix}").status_code == 404
    assert client.get(f"/api/forecasts/february/{suffix}?turbine_id=1").status_code == 404
    assert client.get(f"/api/forecasts/february/{suffix}?turbine_id=3").status_code == 422
    monkeypatch.setattr(
        "app.agent.fetch_archived_weather", lambda *_a, **_kw: february_weather.iloc[:24]
    )
    create_run(issue="2026-01-31")
    assert client.get(f"/api/forecasts/february/{suffix}?turbine_id=1").status_code == 409


@pytest.mark.parametrize("turbine", [1, 2])
def test_real_february_report_exports_672_hours_and_all_agent_logs(turbine):
    response = client.post(f"/api/forecasts/february?turbine_id={turbine}&data_mode=offline")
    assert response.status_code == 200
    assert response.json()["data_source"] == "bundled_archive"
    base = "/api/forecasts/february"
    workbook = xlsx_book(client.get(base + f"/export.xlsx?turbine_id={turbine}"))
    assert workbook["Hourly forecast"].max_row == 673
    assert workbook["Agent log"].max_row == 28 * 7 + 1
    assert workbook["Hourly forecast"].cell(2, 1).value == "2026-02-01T00:00:00+05:00"
    assert workbook["Hourly forecast"].cell(673, 1).value == "2026-02-28T23:00:00+05:00"
    text = pdf_text(client.get(base + f"/report.pdf?turbine_id={turbine}"))
    hourly = text.split("Почасовой прогноз", 1)[1]
    assert len(re.findall(r"\d{2}\.02\.2026 \d{2}:\d{2}", hourly)) == 672
    assert "сохранённый архив" in text
    assert "Синтетические" not in text
    metadata = dict(workbook["Metadata"].iter_rows(min_row=2, values_only=True))
    assert metadata["2026-01-31.weather.manifest_verified"] == "True"


def test_exports_redact_secrets_paths_and_neutralize_excel_formulas():
    run = create_run()
    result = get_forecast(run["run_id"])
    private_text = r"C:\Users\private_user\secret.csv /home/private_user/key.txt OPENAI_API_KEY=hidden-api-value sk-secretstring987 Bearer bearer-secret-token"
    result.quality_warnings.append(private_text)
    result.steps[0].detail = '=HYPERLINK("https://example.com","click")'
    result.steps[1].detail = private_text
    result.weather_provenance["private"] = "private-metadata-secret"
    result.model_artifacts.append({"path": private_text, "api_key": "private-artifact-secret"})
    save_forecast(result)
    text = pdf_text(client.get(f"/api/forecasts/{run['run_id']}/report.pdf"))
    response = client.get(f"/api/forecasts/{run['run_id']}/export.xlsx")
    workbook = xlsx_book(response)
    assert workbook["Agent log"].cell(2, 9).data_type == "s"
    assert workbook["Agent log"].cell(2, 9).value.startswith("=HYPERLINK")
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        text += "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    for secret in (
        "private_user",
        "hidden-api-value",
        "secretstring987",
        "bearer-secret-token",
        "private-metadata-secret",
        "private-artifact-secret",
    ):
        assert secret not in text


def test_historical_report_does_not_claim_metrics_from_different_model():
    run = create_run()
    result = get_forecast(run["run_id"])
    result.model_version = "historical-model"
    save_forecast(result)
    book = xlsx_book(client.get(f"/api/forecasts/{run['run_id']}/export.xlsx"))
    assert book["Model metrics"].max_row == 1
