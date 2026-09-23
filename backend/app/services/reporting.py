"""Local, deterministic forecast reports. No external services or language models."""

import io
import json
import math
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from functools import lru_cache
from statistics import mean
from threading import Lock
from xml.sax.saxutils import escape

from app.catalog import LOCAL_TIMEZONE
from app.config import BACKEND_DIR, settings
from app.schemas import ForecastAccepted

PDF_TYPE = "application/pdf"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NAMES = ["Summary", "Hourly forecast", "Agent log", "Model metrics", "Metadata"]
SOURCE_LABELS = {
    "open_meteo": "Open-Meteo · сеть",
    "cache": "Open-Meteo · локальный кэш",
    "bundled_archive": "Open-Meteo · сохранённый архив",
    "synthetic_test_fixture": "Синтетические данные теста",
}
_FONT_LOCK = Lock()


def safe_text(value) -> str:
    """Defense in depth for text fields; arbitrary metadata is never exported."""
    text = "" if value is None else str(value)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"(?i)(?<!\w)(?:[a-z]:[\\/]|\\\\)[^\s<>\"']+", "[путь скрыт]", text)
    text = re.sub(r"(?<![:\w])/(?:home|Users|tmp|var|etc|opt|mnt)/[^\s<>]+", "[путь скрыт]", text)
    text = re.sub(r"\b(?:sk|nvapi)-[A-Za-z0-9_-]+", "[ключ скрыт]", text)
    text = re.sub(
        r"(?i)\b(?:[a-z_]*api[_-]?key|[a-z_]*token|password|secret|authorization)"
        r"\s*[:=]\s*[^\s,;]+",
        "[секрет скрыт]",
        text,
    )
    text = re.sub(r"(?i)\bBearer\s+[^\s,;]+", "[секрет скрыт]", text)
    return text[:16000]


def local_time(value: datetime) -> str:
    return value.astimezone(LOCAL_TIMEZONE).strftime("%d.%m.%Y %H:%M")


@dataclass
class Report:
    runs: list[ForecastAccepted]
    monthly: bool = False

    def __post_init__(self):
        self.runs = sorted(self.runs, key=lambda run: run.issue_date)
        if not self.runs:
            raise ValueError("Нет данных для отчёта")
        if len({run.turbine_id for run in self.runs}) != 1:
            raise ValueError("Отчёт должен содержать одну турбину")
        if len({run.model_version for run in self.runs}) != 1:
            raise ValueError("Прогнозы используют разные версии модели; пересчитайте месяц")
        for run in self.runs:
            start = datetime.combine(
                run.issue_date + timedelta(days=1), datetime.min.time(), LOCAL_TIMEZONE
            )
            expected = [start + timedelta(hours=i) for i in range(run.horizon_hours)]
            if [point.timestamp for point in run.forecast] != expected:
                raise ValueError("Прогноз содержит неполную или некорректную почасовую сетку")
        if self.monthly:
            expected_dates = [date(2026, 1, 31) + timedelta(days=i) for i in range(28)]
            if [run.issue_date for run in self.runs] != expected_dates or any(
                run.horizon_hours != 24 for run in self.runs
            ):
                raise ValueError("Февральский отчёт требует все 28 выпусков и 672 часа")
        elif len(self.runs) != 1:
            raise ValueError("Ожидался один выпуск прогноза")

    @property
    def points(self):
        return [point for run in self.runs for point in run.forecast]

    @property
    def title(self):
        return "Прогноз на февраль 2026" if self.monthly else "Отчёт о прогнозе мощности"

    @property
    def sources(self):
        return "; ".join(
            sorted(
                {
                    SOURCE_LABELS.get(run.data_source, safe_text(run.data_source))
                    for run in self.runs
                }
            )
        )

    @property
    def warnings(self):
        return sorted({safe_text(warning) for run in self.runs for warning in run.quality_warnings})

    @property
    def peak_hours(self):
        maximum = max(point.normalized_power for point in self.points)
        return [
            point.timestamp
            for point in self.points
            if math.isclose(point.normalized_power, maximum, abs_tol=1e-10)
        ]

    @property
    def summary(self):
        points = self.points
        peak = self.peak_hours
        peak_text = ", ".join(local_time(value) for value in peak[:6])
        if len(peak) > 6:
            peak_text += f"; всего {len(peak)} часов (отмечены в Excel)"
        return (
            f"Для турбины {self.runs[0].turbine_id} рассчитано {len(points)} часов. "
            f"Средняя ожидаемая мощность — {mean(p.normalized_power for p in points):.1%}, "
            f"максимальная — {max(p.normalized_power for p in points):.1%}. "
            f"Часы максимума (UTC+5): {peak_text}. "
            f"Средний прогноз ветра на высоте 100 м — {mean(p.wind_speed_100m_ms for p in points):.1f} м/с; "
            f"температура от {min(p.temperature_c for p in points):.1f} "
            f"до {max(p.temperature_c for p in points):.1f} °C."
        )

    def summary_rows(self):
        points = self.points
        return [
            ("Продукт", "WindFlow AI"),
            ("Отчёт", self.title),
            ("Турбина", self.runs[0].turbine_id),
            ("Начало периода, UTC+5", local_time(points[0].timestamp)),
            ("Конец периода, UTC+5", local_time(points[-1].timestamp)),
            ("Часов", len(points)),
            ("Выпусков прогноза", len(self.runs)),
            ("Средняя мощность, доля 0–1", mean(p.normalized_power for p in points)),
            ("Максимальная мощность, доля 0–1", max(p.normalized_power for p in points)),
            ("Часов максимума", len(self.peak_hours)),
            ("Нормированная выработка, турбино-часы", sum(p.normalized_power for p in points)),
            ("Средний ветер 100 м, м/с", mean(p.wind_speed_100m_ms for p in points)),
            ("Максимальный ветер 100 м, м/с", max(p.wind_speed_100m_ms for p in points)),
            ("Средняя температура, °C", mean(p.temperature_c for p in points)),
            ("Источник погоды", self.sources),
            ("Модель", safe_text(self.runs[0].model_version)),
            ("Резюме", self.summary),
            (
                "Единицы",
                "Мощность нормирована 0–1; выработка в турбино-часах, не кВт·ч. Номинальная мощность не задана.",
            ),
            (
                "Диапазон ошибки",
                "Эмпирический диапазон ошибки; не гарантированный доверительный интервал.",
            ),
            *[("Предупреждение", warning) for warning in self.warnings],
        ]

    def metric_rows(self):
        path = settings.model_dir / "metrics.json"
        if not path.exists():
            return []
        metrics = json.loads(path.read_text(encoding="utf-8"))
        # Never attach metrics of a newer model to a historical saved run.
        if metrics.get("_meta", {}).get("model_version") != self.runs[0].model_version:
            return []
        turbine = metrics.get(f"turbine_{self.runs[0].turbine_id}", {})
        leads = (1,) if self.monthly or self.runs[0].horizon_hours == 24 else (1, 2)
        return [
            (
                self.runs[0].turbine_id,
                "1–24" if lead == 1 else "25–48",
                row.get("mae"),
                row.get("rmse"),
                row.get("r2"),
                row.get("samples"),
                safe_text(row.get("validation_start")),
                safe_text(row.get("validation_end")),
                row.get("baselines", {}).get("wind_curve", {}).get("mae"),
                row.get("interval", {}).get("coverage"),
            )
            for lead in leads
            if (row := turbine.get(f"day_{lead}"))
        ]

    def metadata_rows(self):
        rows = [
            ("product", "WindFlow AI"),
            ("report_generated_at", datetime.now(UTC).isoformat()),
            ("timezone", "Asia/Almaty (UTC+5)"),
            ("model_version", safe_text(self.runs[0].model_version)),
            ("weather_api", "https://previous-runs-api.open-meteo.com/v1/forecast"),
            ("archive_kind", "fixed_lead_offsets"),
            ("publication_time_verified", False),
            ("quality_period", "January 2026 development validation; not independent test"),
            ("power_unit", "normalized 0–1"),
            ("energy_unit", "normalized turbine-hours, not kWh"),
            ("interval_type", "empirical error range, not guaranteed confidence interval"),
        ]
        for run in self.runs:
            prefix = run.issue_date.isoformat()
            rows += [
                (f"{prefix}.run_id", safe_text(run.run_id)),
                (f"{prefix}.issue_at", run.issue_at.isoformat() if run.issue_at else prefix),
                (f"{prefix}.created_at", run.created_at.isoformat()),
                (f"{prefix}.weather_source", safe_text(run.data_source)),
                (f"{prefix}.timing_note", safe_text(run.timing_note)),
            ]
            for key in ("snapshot_sha256", "snapshot_downloaded_at", "manifest_verified"):
                if key in run.weather_provenance:
                    rows.append((f"{prefix}.weather.{key}", safe_text(run.weather_provenance[key])))
            for i, artifact in enumerate(run.model_artifacts):
                for key in ("sha256", "trained_through", "available_from", "role", "lead_days"):
                    if key in artifact:
                        rows.append((f"{prefix}.model_{i + 1}.{key}", safe_text(artifact[key])))
        return rows


def export_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "WindFlow AI"
    workbook.properties.title = report.title
    tables = {
        "Summary": [("Показатель", "Значение"), *report.summary_rows()],
        "Hourly forecast": [
            (
                "Timestamp (UTC+5)",
                "Issue date",
                "Turbine",
                "Power (0–1)",
                "Error range low (0–1)",
                "Error range high (0–1)",
                "Wind 100m (m/s)",
                "Temperature (°C)",
                "Peak hour",
            )
        ],
        "Agent log": [
            (
                "Issue date",
                "Run ID",
                "Step",
                "Stage",
                "Status",
                "Started at",
                "Finished at",
                "Duration (ms)",
                "Detail",
            )
        ],
        "Model metrics": [
            (
                "Turbine",
                "Forecast hours",
                "MAE (0–1)",
                "RMSE (0–1)",
                "R²",
                "Samples",
                "Validation start",
                "Validation end",
                "Baseline MAE",
                "Error range coverage",
            ),
            *report.metric_rows(),
        ],
        "Metadata": [("Key", "Value"), *report.metadata_rows()],
    }
    peak = set(report.peak_hours)
    for run in report.runs:
        for point in run.forecast:
            tables["Hourly forecast"].append(
                (
                    point.timestamp.isoformat(),
                    run.issue_date.isoformat(),
                    run.turbine_id,
                    point.normalized_power,
                    point.confidence_low,
                    point.confidence_high,
                    point.wind_speed_100m_ms,
                    point.temperature_c,
                    point.timestamp in peak,
                )
            )
        for step in run.steps:
            tables["Agent log"].append(
                (
                    run.issue_date.isoformat(),
                    run.run_id,
                    step.id,
                    step.title,
                    step.status,
                    step.started_at.isoformat() if step.started_at else "",
                    step.finished_at.isoformat() if step.finished_at else "",
                    step.duration_ms,
                    step.detail,
                )
            )
    for name in SHEET_NAMES:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for values in tables[name]:
            sheet.append(
                [safe_text(value) if isinstance(value, str) else value for value in values]
            )
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Literal text, including formula-looking messages.
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0000"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163047")
        sheet.row_dimensions[1].height = 32
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 24
        if name in ("Summary", "Metadata"):
            sheet.column_dimensions["A"].width = 48
            sheet.column_dimensions["B"].width = 100
            for row in sheet.iter_rows(min_row=2):
                sheet.row_dimensions[row[0].row].height = max(
                    20, 15 * math.ceil(len(str(row[1].value or "")) / 85)
                )
        if name == "Hourly forecast":
            sheet.column_dimensions["A"].width = 30
        if name == "Agent log":
            sheet.column_dimensions["B"].width = 40
            sheet.column_dimensions["D"].width = 44
            sheet.column_dimensions["I"].width = 95
            for row in sheet.iter_rows(min_row=2):
                sheet.row_dimensions[row[0].row].height = max(
                    30, 15 * math.ceil(len(str(row[8].value or "")) / 80)
                )
    chart = LineChart()
    chart.title = "Прогноз мощности"
    chart.y_axis.title = "Нормированная мощность (0–1)"
    chart.x_axis.title = "Час прогноза (UTC+5)"
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 1
    chart.height, chart.width = 10, 23
    hourly = workbook["Hourly forecast"]
    chart.add_data(
        Reference(hourly, min_col=4, min_row=1, max_row=hourly.max_row), titles_from_data=True
    )
    chart.set_categories(Reference(hourly, min_col=1, min_row=2, max_row=hourly.max_row))
    chart.x_axis.tickLblSkip = max(1, len(report.points) // 8)
    workbook["Summary"].add_chart(chart, f"A{workbook['Summary'].max_row + 3}")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@lru_cache(maxsize=1)
def _register_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    with _FONT_LOCK:
        font_dir = BACKEND_DIR / "assets" / "fonts"
        pdfmetrics.registerFont(TTFont("WindFlow", str(font_dir / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("WindFlowBold", str(font_dir / "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("WindFlow", normal="WindFlow", bold="WindFlowBold")


def _power_chart(report: Report, width=510, height=180):
    from reportlab.graphics.shapes import Drawing, Line, Polygon, PolyLine, String
    from reportlab.lib.colors import HexColor

    drawing = Drawing(width, height)
    x0, y0, plot_w, plot_h = 35, 34, width - 45, height - 54
    points = report.points

    def xy(index, value):
        return (x0 + index / max(1, len(points) - 1) * plot_w, y0 + value * plot_h)

    for level in (0, 0.25, 0.5, 0.75, 1):
        y = y0 + level * plot_h
        drawing.add(Line(x0, y, x0 + plot_w, y, strokeColor=HexColor("#dce4ec")))
        drawing.add(String(3, y - 3, f"{level:.0%}", fontName="WindFlow", fontSize=8))
    band = [coordinate for i, p in enumerate(points) for coordinate in xy(i, p.confidence_high)]
    band += [
        coordinate
        for i in range(len(points) - 1, -1, -1)
        for coordinate in xy(i, points[i].confidence_low)
    ]
    drawing.add(Polygon(band, fillColor=HexColor("#e1eefb"), strokeColor=None))
    line = [coordinate for i, p in enumerate(points) for coordinate in xy(i, p.normalized_power)]
    drawing.add(PolyLine(line, strokeColor=HexColor("#007dba"), strokeWidth=1.5))
    for i in sorted({round((len(points) - 1) * fraction / 4) for fraction in range(5)}):
        x, _ = xy(i, 0)
        label = points[i].timestamp.strftime("%d.%m %H:%M")
        drawing.add(
            String(
                x,
                17,
                label,
                textAnchor="start" if i == 0 else "end" if i == len(points) - 1 else "middle",
                fontName="WindFlow",
                fontSize=7,
            )
        )
    drawing.add(
        String(
            x0,
            height - 7,
            "Мощность и эмпирический диапазон ошибки · UTC+5",
            fontName="WindFlow",
            fontSize=8,
        )
    )
    return drawing


def export_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        CondPageBreak,
        KeepTogether,
        LongTable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _register_fonts()
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=48,
        bottomMargin=42,
        title=report.title,
        author="WindFlow AI",
    )
    body = ParagraphStyle(
        "body",
        fontName="WindFlow",
        fontSize=9,
        leading=14,
        textColor=colors.HexColor("#25394c"),
        spaceAfter=7,
        alignment=TA_LEFT,
    )
    small = ParagraphStyle("small", parent=body, fontSize=7, leading=10)
    heading = ParagraphStyle(
        "heading",
        parent=body,
        fontName="WindFlowBold",
        fontSize=13,
        leading=18,
        spaceBefore=12,
        keepWithNext=True,
    )
    title = ParagraphStyle("title", parent=heading, fontSize=23, leading=29, spaceAfter=12)

    def paragraph(text, style=body):
        return Paragraph(escape(safe_text(text)).replace("\n", "<br/>"), style)

    def table_style(header=True):
        commands = [
            ("FONTNAME", (0, 0), (-1, -1), "WindFlow"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            (
                "ROWBACKGROUNDS",
                (0, 1 if header else 0),
                (-1, -1),
                [colors.white, colors.HexColor("#f0f5f9")],
            ),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#007dba")),
        ]
        if header:
            commands += [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcebf5")),
                ("FONTNAME", (0, 0), (-1, 0), "WindFlowBold"),
            ]
        return TableStyle(commands)

    first = report.runs[0]
    points = report.points
    story = [paragraph("WINDFLOW AI", small), paragraph(report.title, title)]
    parameters = [
        ("Турбина", str(first.turbine_id)),
        (
            "Период (UTC+5)",
            f"{local_time(points[0].timestamp)} — {local_time(points[-1].timestamp)}",
        ),
        (
            "Выпуск прогноза",
            f"{first.issue_date:%d.%m.%Y} — {report.runs[-1].issue_date:%d.%m.%Y}"
            if report.monthly
            else local_time(first.issue_at)
            if first.issue_at
            else first.issue_date.isoformat(),
        ),
        ("Объём", f"{len(points)} часов · {len(report.runs)} выпусков"),
        ("Модель", first.model_version),
        ("Погода", report.sources),
    ]
    table = Table(
        [[paragraph(k, small), paragraph(v, small)] for k, v in parameters], colWidths=[125, 390]
    )
    table.setStyle(table_style(False))
    table.setStyle(
        TableStyle(
            [
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story += [
        table,
        paragraph("Краткое резюме", heading),
        paragraph(report.summary),
        _power_chart(report),
    ]
    story += [
        paragraph(
            f"Нормированная выработка: {sum(p.normalized_power for p in points):.2f} турбино-часов. Мощность показана в процентах нормированной шкалы. Перевод в кВт·ч требует номинальной мощности турбины.",
            small,
        )
    ]
    story += [
        paragraph("Как читать результат", heading),
        paragraph(
            "Диапазон на графике получен из ошибок модели на проверочном периоде. Он не гарантирует покрытие будущих значений. Прогноз предназначен для планирования; фактическая выработка может отличаться.",
            small,
        ),
    ]
    story += [
        paragraph("Источник и предупреждения", heading),
        paragraph(
            "Источник: Open-Meteo Previous Runs, архив прогнозов с фиксированным упреждением 24/48 часов. Фактическая будущая погода не используется. Точное время публикации каждого погодного выпуска архивом не подтверждено.",
            small,
        ),
    ]
    story += [paragraph(warning, small) for warning in report.warnings]
    metrics = report.metric_rows()
    story += [paragraph("Качество модели", heading)]
    if metrics:
        rows = [["Часы", "MAE", "RMSE", "R²", "Часов проверки"]]
        rows += [
            [row[1], f"{row[2]:.4f}", f"{row[3]:.4f}", f"{row[4]:.4f}", row[5]] for row in metrics
        ]
        table = Table(rows, colWidths=[75, 95, 95, 95, 155])
        table.setStyle(table_style())
        story += [
            KeepTogether(
                [
                    table,
                    Spacer(1, 6),
                    paragraph(
                        "Оценка на январе 2026: период проверки при разработке, не независимый закрытый тест. MAE и RMSE относятся к шкале мощности 0–1. Метрики часов 25–48 рассчитаны отдельно от часов 1–24.",
                        small,
                    ),
                ]
            )
        ]
    else:
        story.append(paragraph("Метрики для сохранённой версии модели недоступны.", small))
    story += [
        CondPageBreak(180),
        paragraph("Почасовой прогноз", heading),
        paragraph(
            "Время — UTC+5; мощность и диапазон — проценты нормированной мощности; ветер на высоте 100 м.",
            small,
        ),
    ]
    hourly = [
        [
            paragraph(text, small)
            for text in ("Дата и час", "Мощность, %", "Диапазон, %", "Ветер, м/с", "Темп., °C")
        ]
    ]
    for point in points:
        hourly.append(
            [
                local_time(point.timestamp),
                f"{point.normalized_power * 100:.2f}",
                f"{point.confidence_low * 100:.1f} – {point.confidence_high * 100:.1f}",
                f"{point.wind_speed_100m_ms:.2f}",
                f"{point.temperature_c:.1f}",
            ]
        )
    table = LongTable(hourly, colWidths=[130, 95, 120, 90, 80], repeatRows=1)
    table.setStyle(table_style())
    table.setStyle(
        TableStyle(
            [
                ("TOPPADDING", (0, 1), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 3),
            ]
        )
    )
    story.append(table)

    def page_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("WindFlow", 7)
        canvas.setFillColor(colors.HexColor("#617384"))
        canvas.drawString(
            40, 24, f"WindFlow AI · Турбина {first.turbine_id} · {safe_text(first.model_version)}"
        )
        canvas.drawRightString(A4[0] - 40, 24, f"Страница {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()
